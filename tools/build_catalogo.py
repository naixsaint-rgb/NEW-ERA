import json, sys, os, glob, re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

sp = sys.argv[1]
OUT = f'{sp}/catalogo/out'
os.makedirs(OUT, exist_ok=True)
ORDER = ['cobre_metales','psicodelicos','nicotina_colinergicos','infecciones','metabolico_cetonas','luz_energia','plantas_tradiciones','protocolos_heterodoxos','hormonas_peptidos','sangre_plasma_quelacion','cuerpo_mente','ambiente_toxicos','nutricion_alimentos','farmacos_reposicionados','relatos_familias']
files = {os.path.basename(p)[:-5]: p for p in glob.glob(f'{sp}/catalogo/verified/*.json')}
for p in glob.glob(f'{sp}/catalogo/research/*.json'):
    k = os.path.basename(p)[:-5]
    files.setdefault(k, p)  # fall back to unverified research if verification failed
keys = [k for k in ORDER if k in files] + sorted(k for k in files if k not in ORDER)
import os as _os
themes = []
for k in keys:
    try:
        d = json.load(open(files[k]))
        d['_src'] = 'verified' if '/verified/' in files[k] else 'research'
        gp = f'{sp}/catalogo/grupos/{k}.json'
        grupos = None
        if _os.path.exists(gp):
            try:
                g = json.load(open(gp)); grupos = g.get('grupos') if isinstance(g, dict) else g
            except Exception as e: print('grupos ERR', k, e)
        rows_ = d.get('rows', [])
        if grupos and len(grupos) == len(rows_):
            for row, gname in zip(rows_, grupos): row['grupo'] = (gname or row.get('intervencion') or '').strip()
        else:
            if grupos is not None: print('grupos: longitud no coincide en', k, len(grupos), len(rows_))
            for row in rows_: row['grupo'] = (row.get('intervencion') or '').strip()
        themes.append(d)
    except Exception as e:
        print('ERROR leyendo', k, e)

COLS = [('tema','Tema',14),('grupo','Intervención (grupo)',24),('intervencion','Intervención (tal como la nombra la fuente)',26),('tradicion_o_angulo','Tradición o ángulo',18),('estudio','Estudio / fuente',34),('tipo_fuente','Tipo de fuente',16),('diseno','Diseño',22),('n','n',8),('poblacion','Población',18),('duracion','Duración',12),('dosis_o_protocolo','Dosis o protocolo',28),('que_midio','Qué midió',24),('resultado_reportado','Resultado reportado (cifra tal cual)',44),('direccion_segun_la_fuente','Dirección según la fuente',14),('efectos_adversos_reportados','Efectos adversos reportados',26),('financiacion_y_conflictos','Financiación y conflictos',28),('como_probarlo_en_casa','Cómo probarlo en casa (según la fuente)',30),('coste_aprox','Coste aprox.',12),('verificacion','Verificación documental',26),('identificador','Identificador',18),('enlace','Enlace',40)]
FONT = Font(name='Arial', size=10)
BOLD = Font(name='Arial', size=10, bold=True)
HFILL = PatternFill('solid', fgColor='DDDDDD')
WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()
ws = wb.active; ws.title = 'Filas'
ws.append([c[1] for c in COLS])
for i, c in enumerate(COLS, 1):
    cell = ws.cell(row=1, column=i); cell.font = BOLD; cell.fill = HFILL; cell.alignment = WRAP
    ws.column_dimensions[get_column_letter(i)].width = c[2]
nrows = 0
for t in themes:
    for r in t.get('rows', []):
        row = []
        for key, _, _ in COLS:
            v = t['key'] if key == 'tema' else r.get(key, '')
            if isinstance(v, (list, dict)): v = json.dumps(v, ensure_ascii=False)
            row.append(v if v is not None else '')
        ws.append(row); nrows += 1
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row: cell.font = FONT; cell.alignment = WRAP
ws.freeze_panes = 'C2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{ws.max_row}'
# column letters for formulas
L = {key: get_column_letter(i) for i, (key, _, _) in enumerate(COLS, 1)}
last = ws.max_row

# Resumen por intervención (formulas, not hardcoded)
rs = wb.create_sheet('Resumen')
heads = ['Tema','Intervención (grupo)','Nº estudios o fuentes','Mejora','Sin cambio','Empeora','Mixto','Citas encontradas','Citas no encontradas']
rs.append(heads)
for i, h in enumerate(heads, 1):
    c = rs.cell(row=1, column=i); c.font = BOLD; c.fill = HFILL; c.alignment = WRAP
widths = [16, 34, 12, 10, 10, 10, 10, 14, 14]
for i, w in enumerate(widths, 1): rs.column_dimensions[get_column_letter(i)].width = w
seen = []
for t in themes:
    for r in t.get('rows', []):
        pair = (t['key'], (r.get('grupo') or r.get('intervencion') or '').strip())
        if pair[1] and pair not in seen: seen.append(pair)
rng = lambda key: f"Filas!${L[key]}$2:${L[key]}${last}"
for idx, (tema, interv) in enumerate(seen, 2):
    rs.cell(row=idx, column=1, value=tema).font = FONT
    rs.cell(row=idx, column=2, value=interv).font = FONT
    base = f'{rng("tema")},$A{idx},{rng("grupo")},$B{idx}'
    rs.cell(row=idx, column=3, value=f'=COUNTIFS({base})').font = FONT
    for col, lab in zip([4,5,6,7], ['mejora','sin cambio','empeora','mixto']):
        rs.cell(row=idx, column=col, value=f'=COUNTIFS({base},{rng("direccion_segun_la_fuente")},"{lab}")').font = FONT
    rs.cell(row=idx, column=8, value=f'=COUNTIFS({base},{rng("verificacion")},"cita encontrada*")+COUNTIFS({base},{rng("verificacion")},"página existe*")').font = FONT
    rs.cell(row=idx, column=9, value=f'=COUNTIFS({base},{rng("verificacion")},"cita no encontrada*")+COUNTIFS({base},{rng("verificacion")},"página no accesible*")').font = FONT
tot = len(seen) + 2
rs.cell(row=tot, column=2, value='TOTAL').font = BOLD
for col in range(3, 10):
    cl = get_column_letter(col)
    rs.cell(row=tot, column=col, value=f'=SUM({cl}2:{cl}{tot-1})').font = BOLD
rs.freeze_panes = 'C2'
rs.auto_filter.ref = f'A1:I{tot-1}'

# Dosieres
ds = wb.create_sheet('Dosieres')
dh = ['Tema','Título','Tesis','Defensores','Mecanismos propuestos','Consenso y quién lo paga','Lo enterrado','Protocolos reportados','Riesgos reconocidos por defensores','Fuentes heterodoxas','Ideas para la app','Notas del investigador']
ds.append(dh)
for i, h in enumerate(dh, 1):
    c = ds.cell(row=1, column=i); c.font = BOLD; c.fill = HFILL; c.alignment = WRAP
    ds.column_dimensions[get_column_letter(i)].width = 16 if i <= 2 else 45
def J(v): return '\n'.join(f'• {x}' for x in v) if isinstance(v, list) else (v or '')
for t in themes:
    d = t.get('dosier', {})
    ds.append([t['key'], t.get('title',''), d.get('tesis',''), J(d.get('defensores')), J(d.get('mecanismos')), d.get('consenso_y_quien_lo_paga',''), J(d.get('lo_enterrado')), J(d.get('protocolos_reportados')), J(d.get('riesgos_reconocidos_por_defensores')), J(d.get('fuentes_heterodoxas')), J(d.get('ideas_para_la_app')), t.get('notas','')])
for row in ds.iter_rows(min_row=2, max_row=ds.max_row):
    for cell in row: cell.font = FONT; cell.alignment = WRAP
ds.freeze_panes = 'C2'

# Enterrado
es = wb.create_sheet('Enterrado')
es.append(['Tema','Qué se enterró o quedó sin publicar','Fuente'])
for i in range(1, 4):
    c = es.cell(row=1, column=i); c.font = BOLD; c.fill = HFILL
es.column_dimensions['A'].width = 18; es.column_dimensions['B'].width = 80; es.column_dimensions['C'].width = 30
for t in themes:
    for x in t.get('dosier', {}).get('lo_enterrado', []):
        es.append([t['key'], x, 'dosier'])
    for r in t.get('rows', []):
        if r.get('tipo_fuente') == 'registro de ensayo':
            es.append([t['key'], f"{r.get('intervencion','')} · {r.get('estudio','')} · {r.get('resultado_reportado','')}", r.get('enlace','')])
for row in es.iter_rows(min_row=2, max_row=es.max_row):
    for cell in row: cell.font = FONT; cell.alignment = WRAP

# Leyenda
ls = wb.create_sheet('Leyenda', 0)
ls.column_dimensions['A'].width = 120
lines = [
 'CATÁLOGO DE INTERVENCIONES PARA ALZHEIMER · investigado con postura contraria al consenso · 12/10/2026',
 '',
 'Cómo se hizo: 15 temas, cada uno investigado por un agente con la instrucción de no fiarse del consenso, seguir el dinero, buscar donde el consenso no mira y no descartar nada por pequeño, antiguo, animal o heterodoxo. Fuentes: Europe PMC y PubMed por API, ClinicalTrials.gov, y webs, foros y libros heterodoxos.',
 'Nadie hace de juez. Las cifras se transcriben tal cual las da cada fuente. La columna "Dirección según la fuente" es lo que dice la propia fuente, no una opinión nuestra.',
 'Verificación documental: un segundo agente solo comprobó que la cita existe y que la cifra aparece en la fuente. No borró ni cambió nada; si una cifra difiere, lo anota en la columna "Verificación documental" y tú decides.',
 'Financiación y conflictos: lo que declara cada estudio; "no consta" significa que la fuente no lo dice, no que no exista.',
 'Columna Intervención (grupo): nombre canónico asignado por un agente para poder sumar filas de la misma intervención; la columna siguiente conserva el nombre tal como lo da la fuente. Hoja Filas: una fila por estudio o fuente. Hoja Resumen: cuenta por intervención con fórmulas sobre la hoja Filas (se recalcula sola si editas Filas). Hoja Dosieres: la tesis heterodoxa de cada tema, sus defensores, lo enterrado y las ideas para la app. Hoja Enterrado: ensayos sin resultados publicados y líneas abandonadas.',
 'Las conclusiones las saca la familia. Este archivo no es consejo médico; es material de investigación.',
 f'Filas: {nrows}. Temas: {len(themes)}. Intervenciones distintas: {len(seen)}.',
]
for i, l in enumerate(lines, 1):
    c = ls.cell(row=i, column=1, value=l); c.font = BOLD if i == 1 else FONT; c.alignment = WRAP
wb.save(f'{OUT}/CATALOGO.xlsx')

# Markdown for Obsidian
md = ['# NACHO · Catálogo de intervenciones para Alzheimer (postura contraria al consenso)', '', f'12/10/2026 · {nrows} filas · {len(themes)} temas · {len(seen)} intervenciones distintas. Una fila por estudio o fuente; cifras tal cual las da la fuente; nadie hace de juez. Verificación documental = la cita existe y la cifra aparece. Las conclusiones las saca la familia.', '']
dos = ['# NACHO · Dosieres por tema (postura contraria al consenso)', '']
for t in themes:
    d = t.get('dosier', {})
    md.append(f"\n## {t.get('title', t['key'])} (`{t['key']}`)\n")
    md.append(f"**Tesis:** {d.get('tesis','')}\n")
    for r in t.get('rows', []):
        parts = [f"**{r.get('grupo') or r.get('intervencion','')}**" + (f" ({r.get('intervencion','')})" if r.get('grupo') and r.get('grupo') != r.get('intervencion') else ''), r.get('estudio',''), r.get('tipo_fuente',''), f"n={r.get('n','')}", r.get('poblacion',''), r.get('duracion','')]
        if r.get('dosis_o_protocolo'): parts.append('protocolo: ' + r['dosis_o_protocolo'])
        parts.append(f"midió {r.get('que_midio','')} → {r.get('resultado_reportado','')} ({r.get('direccion_segun_la_fuente','')})")
        if r.get('efectos_adversos_reportados'): parts.append('adversos: ' + r['efectos_adversos_reportados'])
        parts.append('financiación: ' + (r.get('financiacion_y_conflictos') or 'no consta'))
        if r.get('como_probarlo_en_casa'): parts.append('en casa: ' + r['como_probarlo_en_casa'])
        if r.get('verificacion'): parts.append('verificación: ' + r['verificacion'])
        parts.append(f"[fuente]({r.get('enlace','')})")
        md.append('- ' + ' · '.join(str(p) for p in parts if p))
    dos.append(f"\n## {t.get('title', t['key'])} (`{t['key']}`)\n")
    for lab, key in [('Tesis','tesis'),('Defensores','defensores'),('Mecanismos propuestos','mecanismos'),('Consenso y quién lo paga','consenso_y_quien_lo_paga'),('Lo enterrado','lo_enterrado'),('Protocolos reportados','protocolos_reportados'),('Riesgos reconocidos por los propios defensores','riesgos_reconocidos_por_defensores'),('Fuentes heterodoxas','fuentes_heterodoxas'),('Ideas para la app','ideas_para_la_app')]:
        v = d.get(key)
        if isinstance(v, list): dos.append(f"**{lab}:**\n" + '\n'.join(f'- {x}' for x in v) + '\n')
        else: dos.append(f"**{lab}:** {v or ''}\n")
    if t.get('notas'): dos.append(f"*Notas del investigador:* {t['notas']}\n")
open(f'{OUT}/CATALOGO.md','w').write('\n'.join(md))
open(f'{OUT}/DOSIERES.md','w').write('\n'.join(dos))
print(json.dumps({'temas': [t['key'] + ('' if t['_src']=='verified' else ' (sin verificar)') for t in themes], 'filas': nrows, 'intervenciones': len(seen)}, ensure_ascii=False))
